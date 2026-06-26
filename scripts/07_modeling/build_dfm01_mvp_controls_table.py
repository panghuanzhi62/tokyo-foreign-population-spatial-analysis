"""Complete the DF_M01 MVP controls table for the 251 Greater Tokyo frame.

Joins the three controls already found in the local Tokyo feature tables to the 251
frame, adds a 2024 total-population column (Basic Resident Register) and a 2024
population-density where area is available, backfills rail accessibility for
municipalities present in the wider 227-row tables but absent from v5, and itemizes
every remaining gap. It does NOT download data, does NOT run GIS, does NOT invent a
commercial-density proxy, and does NOT run models.

Outputs:
  LOCAL-ONLY (never committed):
    data_processed_official/model_panel/tokyo_china_dfm01_mvp_controls_joined_local_only.csv
  COMMITTABLE:
    outputs/modeling/dfm01_mvp/dfm01_mvp_controls_coverage_gap.csv
    outputs/modeling/dfm01_mvp/dfm01_mvp_controls_completion_inventory.csv

Usage:
    python build_dfm01_mvp_controls_table.py --project-root E:\\rsch\\laborJapan
"""
import argparse
import csv
import math
from pathlib import Path

import openpyxl


def load_csv_bom(path):
    with open(path, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def fnum(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def classify_gap(code, pref):
    tokyo_islands = {"13361","13362","13363","13364","13381","13382","13401","13402","13421"}
    tokyo_mountain = {"13307","13308"}
    if code in tokyo_islands:
        return "tokyo_island_municipality"
    if code in tokyo_mountain:
        return "tokyo_mountain_village_town"
    if pref == "12":
        return "chiba_boso_or_eastern_city_town_village"
    if pref == "11":
        return "saitama_chichibu_town_village"
    return "other"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--project-root", required=True)
    args = ap.parse_args()
    root = Path(args.project_root)

    # ---- 251 frame (code, prefecture, name) ----
    isa = list(csv.DictReader(open(
        root/"data_processed_official/isa_moj_panel/isa_moj_chinese_stock_t2_2023_12_2025_06.csv",
        encoding="utf-8")))
    frame = {}
    for r in isa:
        frame[r["municipality_code"]] = (r["prefecture_code"], r["municipality_name"])
    codes = sorted(frame)

    # ---- v5 features (primary control source) ----
    v5 = {str(r["N03_007"]).strip().zfill(5): r for r in load_csv_bom(
        root/"data_raw/tokyo_features_v5_extended_with_density.csv")}
    # v2 raw fields
    v2 = {str(r["N03_007"]).strip().zfill(5): r for r in load_csv_bom(
        root/"data_raw/tokyo_features_v2.csv")}
    # pop_with_station: 227 coverage, raw dist_to_station_m (secondary rail source)
    popst = {str(r["N03_007"]).strip().zfill(5): r for r in load_csv_bom(
        root/"data_raw/tokyo_pop_with_station_features.csv")}

    # ---- 2024 total population (Basic Resident Register, 000959256.xlsx) ----
    wb = openpyxl.load_workbook(root/"data_raw/000959256.xlsx", read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    pop2024 = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 7:
            continue
        code6 = row[0]
        if code6 in (None, "-", ""):
            continue
        code6 = str(code6).strip()
        if len(code6) == 6 and code6.isdigit():
            pop2024[code6[:5]] = fnum(row[5])
    wb.close()

    rows = []
    gap_rows = []
    cov = {"rail": 0, "housing": 0, "popdens2020": 0, "popdens2024": 0,
           "pop2024": 0, "rail_backfill": 0, "commercial": 0, "in_v5": 0}

    for code in codes:
        pref, name = frame[code]
        in_v5 = code in v5
        if in_v5:
            cov["in_v5"] += 1
        r5 = v5.get(code, {})
        r2 = v2.get(code, {})
        rps = popst.get(code, {})

        rail = fnum(r5.get("log_dist_to_station_m"))
        rail_src = "v5_log_dist_to_station_m" if rail is not None else ""
        rail_backfilled = "NO"
        # backfill rail from 227-row pop_with_station table if v5 lacks it
        if rail is None:
            dst = fnum(rps.get("dist_to_station_m"))
            if dst is not None and dst > 0:
                rail = math.log(dst)
                rail_src = "backfill_pop_with_station_log(dist_to_station_m)"
                rail_backfilled = "YES"
                cov["rail_backfill"] += 1

        housing = fnum(r5.get("log_median_land_price_jpy"))
        popdens2020 = fnum(r5.get("population_density"))
        total_pop2020 = fnum(r5.get("total_pop"))
        area = fnum(r5.get("area_sqkm"))
        p2024 = pop2024.get(code)
        popdens2024 = (p2024 / area) if (p2024 is not None and area not in (None, 0)) else None

        dist_raw = fnum(r2.get("dist_to_station_m")) if r2 else fnum(rps.get("dist_to_station_m"))
        landprice_raw = fnum(r2.get("median_land_price_jpy")) if r2 else None

        if rail is not None: cov["rail"] += 1
        if housing is not None: cov["housing"] += 1
        if popdens2020 is not None: cov["popdens2020"] += 1
        if popdens2024 is not None: cov["popdens2024"] += 1
        if p2024 is not None: cov["pop2024"] += 1

        rows.append({
            "municipality_code": code, "municipality_name": name, "prefecture": pref,
            "in_v5_features": "YES" if in_v5 else "NO",
            "rail_accessibility": "" if rail is None else rail,
            "rail_source": rail_src, "rail_backfilled": rail_backfilled,
            "housing_cost": "" if housing is None else housing,
            "population_density": "" if popdens2020 is None else popdens2020,
            "population_density_2024": "" if popdens2024 is None else popdens2024,
            "commercial_density": "",
            "dist_to_station_m": "" if dist_raw is None else dist_raw,
            "median_land_price_jpy": "" if landprice_raw is None else landprice_raw,
            "total_pop": "" if total_pop2020 is None else total_pop2020,
            "area_sqkm": "" if area is None else area,
            "population_2024": "" if p2024 is None else p2024,
        })

        gap_rows.append({
            "municipality_code": code, "municipality_name": name, "prefecture": pref,
            "in_v5_features": "YES" if in_v5 else "NO",
            "rail_accessibility_available": "YES" if rail is not None else "NO",
            "housing_cost_available": "YES" if housing is not None else "NO",
            "population_density_available": "YES" if (popdens2020 is not None or popdens2024 is not None) else "NO",
            "commercial_density_available": "NO",
            "gap_reason": "in_v5_full" if in_v5 else classify_gap(code, pref),
            "recommended_gap_action": ("use_v5" if in_v5
                else ("backfill_rail_only_construct_housing_density_from_official"
                      if rail is not None else "construct_all_from_official_or_gis")),
        })

    # ---- write LOCAL-ONLY joined controls table ----
    panel_dir = root/"data_processed_official/model_panel"
    panel_dir.mkdir(parents=True, exist_ok=True)
    joined_path = panel_dir/"tokyo_china_dfm01_mvp_controls_joined_local_only.csv"
    with open(joined_path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)
    print(f"[local-only] controls table: {joined_path} ({len(rows)} rows)")

    # ---- write COMMITTABLE coverage gap table ----
    out_dir = root/"outputs/modeling/dfm01_mvp"
    out_dir.mkdir(parents=True, exist_ok=True)
    gap_path = out_dir/"dfm01_mvp_controls_coverage_gap.csv"
    with open(gap_path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(gap_rows[0].keys()))
        w.writeheader(); w.writerows(gap_rows)
    print(f"[committable] coverage gap: {gap_path}")

    # ---- write COMMITTABLE per-control completion inventory ----
    n = len(codes)
    inv = [
      {"control":"rail_accessibility","primary_source":"data_raw/tokyo_features_v5 log_dist_to_station_m (+pop_with_station backfill)",
       "field_mapped":"log_dist_to_station_m","temporal_reference":"N02-22 stations ~2022","year_class":"B_baseline_structural",
       "coverage_count":f"{cov['rail']}/{n}","joined":"PARTIAL","ready_for_full_251":"NO",
       "recommended_action":"map_field_and_join_then_fill_remaining","notes":f"{cov['rail_backfill']} municipalities backfilled from 227-row pop_with_station table"},
      {"control":"housing_cost","primary_source":"data_raw/tokyo_features_v5 log_median_land_price_jpy",
       "field_mapped":"log_median_land_price_jpy","temporal_reference":"L01-24 land price 2024","year_class":"A_direct_2024",
       "coverage_count":f"{cov['housing']}/{n}","joined":"PARTIAL","ready_for_full_251":"NO",
       "recommended_action":"map_field_and_join_then_construct_missing_from_official","notes":"raw median_land_price_jpy preserved from v2"},
      {"control":"population_density","primary_source":"data_raw/tokyo_features_v5 population_density (2020); 000959256.xlsx pop_2024 + area",
       "field_mapped":"population_density / population_density_2024","temporal_reference":"2020 census density; 2024 pop where area available","year_class":"B_baseline_structural",
       "coverage_count":f"{cov['popdens2020']}/{n} (2020); {cov['popdens2024']}/{n} (2024 where area)","joined":"PARTIAL","ready_for_full_251":"NO",
       "recommended_action":"join_2020_density_and_or_build_2024_density_needs_area_for_33","notes":f"2024 total population available {cov['pop2024']}/{n}; area only {cov['popdens2024']}/{n}"},
      {"control":"commercial_density","primary_source":"none",
       "field_mapped":"","temporal_reference":"none","year_class":"D_unusable",
       "coverage_count":f"0/{n}","joined":"NO","ready_for_full_251":"NO",
       "recommended_action":"construct_new_source_economic_census","notes":"no establishment-count source locally; do not substitute DF_M04 brokers or POIs"},
    ]
    inv_path = out_dir/"dfm01_mvp_controls_completion_inventory.csv"
    with open(inv_path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(inv[0].keys()))
        w.writeheader(); w.writerows(inv)
    print(f"[committable] completion inventory: {inv_path}")

    print("COVERAGE:", cov)
    missing = [c for c in codes if c not in v5]
    print("v5 coverage:", cov["in_v5"], "/", n, "; missing:", len(missing))


if __name__ == "__main__":
    main()
